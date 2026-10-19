"""Banc d'essai : jeu de référence tiré d'une réponse DOCX (fictive, fabriquée
ici) et mesures sur le corpus « ventes » indexé en mémoire, sans modèle réel."""

import json
import shutil
import sys
from pathlib import Path

import docx
import pytest
from pydantic import ValidationError

from luciole.evaluation.bench import main as bench_main
from luciole.evaluation.bench import noise, summarize
from luciole.evaluation.reference import (
    Expected,
    ReferenceItem,
    parse_response,
    read_reference,
    references,
    write_reference,
)
from luciole.evaluation.spec import BenchSpec, load_bench_spec

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tests"))
from test_answer import EVASIVE, GOOD, ScaledReranker, ScriptedLLM  # noqa: E402
from test_index_pipeline import Setup  # noqa: E402

MRAE = load_bench_spec(ROOT / "configs" / "profiles" / "mrae")


def remark(document, number: int, text: str) -> None:
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = f"Remarque {number}"
    table.cell(1, 0).text = text


@pytest.fixture()
def response(tmp_path):
    """Réponse fictive, à la forme des réponses réelles aux avis MRAe."""
    d = docx.Document()
    d.add_paragraph("Préambule : le projet fictif du Moulin, voir Tome 9 p. 1.")
    remark(d, 1, "L'autorité recommande de préciser les mesures pour les chauves-souris.")
    d.add_paragraph("Voir étude d'impact sur l'environnement : Tome 4 – Milieu naturel :")
    d.add_paragraph("p. 432-433 (impacts bruts) et 445 (mesure ECO-E6)")
    d.add_paragraph("Le choix des variantes est présenté dès la page 59 du Tome 1.")
    d.add_paragraph("Voir Annexe 1 – Note de suivi : p. 40")
    d.add_paragraph("Comme le rappelle le guide national (page 18), rien de plus.")
    remark(d, 2, "L'autorité recommande de reprendre l'analyse paysagère.")
    d.add_paragraph("Le pétitionnaire prend acte et complétera le dossier.")
    remark(d, 3, "L'autorité recommande de compléter les cartes.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Voir Tome 5 – Volet paysage (pages 202-203), p. 331 et 393"
    path = tmp_path / "Moulin_Réponse.docx"
    d.save(path)
    return path


def test_reference_tiree_de_la_reponse(response):
    items = parse_response(response, MRAE.reference, dossier="Dossier Moulin", prefix="Moulin")
    assert [i.id for i in items] == ["Moulin-R1", "Moulin-R2", "Moulin-R3"]
    first, second, third = items
    assert first.question.startswith("L'autorité recommande de préciser")
    assert first.dossier == "Dossier Moulin" and first.source == "Moulin_Réponse.docx"
    # Renvoi au bloc suivant, « et 445 », page avant le tome ; annexes, guides
    # et préambule ignorés.
    assert first.expected == [
        Expected(part="1", pages=[59]),
        Expected(part="4", pages=[432, 433, 445]),
    ]
    assert second.expected == []
    assert third.expected == [Expected(part="5", pages=[202, 203, 331, 393])]


def test_renvois_dans_un_bloc():
    spec = MRAE.reference
    assert references("Tome 4 : p. 475-485", spec) == ({"4": set(range(475, 486))}, None)
    assert references("Tome 4 : p. 10-900", spec) == ({"4": {10}}, None)  # plage aberrante
    assert references("Voir Tome 2 – Milieu physique :", spec) == ({}, "2")
    assert references("p. 7", spec, "2") == ({"2": {7}}, None)
    assert references("p. 7", spec) == ({}, None)


def test_jsonl_aller_retour(tmp_path, response):
    items = parse_response(response, MRAE.reference, dossier="Dossier Moulin", prefix="Moulin")
    path = tmp_path / "ref.jsonl"
    write_reference(items, path)
    assert read_reference(path) == items
    with pytest.raises(ValueError, match="déjà"):
        write_reference(items, path, append=True)
    write_reference([ReferenceItem(id="Autre-R1", dossier="d", question="q")], path, append=True)
    assert len(read_reference(path)) == 4
    path.write_text('{"id": "x", "dossier": "d"}\n', encoding="utf-8")
    with pytest.raises(ValueError, match="ligne 1"):
        read_reference(path)


def test_reglages_controles():
    raw = MRAE.model_dump()
    raw["reference"]["page_pattern"] = "(p"
    with pytest.raises(ValidationError, match="illisible"):
        BenchSpec.model_validate(raw)
    raw = MRAE.model_dump()
    raw["metrics"]["k"] = [0]
    with pytest.raises(ValidationError):
        BenchSpec.model_validate(raw)
    with pytest.raises(FileNotFoundError):
        load_bench_spec(ROOT / "configs" / "profiles" / "ventes")


# --- banc d'essai sur le corpus « ventes » ----------------------------------
# La « pièce » d'un document de vente est son numéro de réunion.
BENCH = """
reference:
  remark_pattern: '^Remarque (\\d+)$'
  part_pattern: 'r[ée]union[-_ ]*(\\d+)'
  page_pattern: '\\bp\\.\\s*(\\d+)'
metrics:
  part_field: numero_reunion
  k: [1, 3]
  page_tolerance: 0
baseline:
  source: essai
"""
QUESTION = "Pourquoi l'atelier trouve douze mille euros au-dessus de ce qu'on peut engager ?"


@pytest.fixture()
def bench(tmp_path):
    s = Setup(tmp_path)
    s.pipeline.run(s.root)
    profiles = tmp_path / "profiles"
    shutil.copytree(ROOT / "configs" / "profiles" / "ventes", profiles / "ventes")
    (profiles / "ventes" / "evaluation").mkdir()
    (profiles / "ventes" / "evaluation" / "bench.yaml").write_text(BENCH, encoding="utf-8")
    reference = tmp_path / "ref.jsonl"
    write_reference(
        [
            ReferenceItem(
                id="Dubois-R1",
                dossier="atelier-dubois",
                question=QUESTION,
                expected=[Expected(part="3", pages=[1])],
            ),
            ReferenceItem(id="Dubois-R2", dossier="atelier-dubois", question="Le budget ?"),
            ReferenceItem(
                id="Roche-R1",
                dossier="cabinet-roche",
                question=QUESTION,
                expected=[Expected(part="3", pages=[1])],  # absent de ce dossier
            ),
        ],
        reference,
    )

    def run(*extra, llm=None):
        lines: list[str] = []
        code = bench_main(
            [
                "--profile",
                "ventes",
                "--profiles-dir",
                str(profiles),
                "--hardware",
                "cpu",
                "--store",
                str(s.store.path),
                "--reference",
                str(reference),
                "--report",
                str(tmp_path / "rapport.json"),
                *extra,
            ],
            log=lines.append,
            embedder_factory=lambda spec, hw: s.embedder,
            reranker_factory=lambda spec, hw: ScaledReranker(),
            qdrant=s.qdrant,
            opensearch=s.opensearch,
            llm_factory=(lambda hw, gen: llm) if llm else None,
        )
        assert code == 0
        return "\n".join(lines), json.loads((tmp_path / "rapport.json").read_text("utf-8"))

    yield run
    s.store.close()


def test_banc_recherche(bench):
    out, rep = bench("--runs", "2")
    summary = rep["runs"][0]["summary"]
    assert summary["remarques"] == 3 and summary["avec_renvoi"] == 2
    assert summary["bon_tome@1"] == 1 and summary["bonne_page@1"] == 1  # Roche n'a pas de R3
    assert summary["mrr_tome"] == 0.5 and "esquives" not in summary
    first = rep["runs"][0]["items"][0]
    assert first["part_rank"] == 1 and first["citations"][0].startswith("atelier-dubois/")
    # Le dossier limite la recherche : aucun passage de Roche pour une remarque Dubois.
    assert all(c.startswith("atelier-dubois/") for c in first["citations"])
    assert rep["noise"] == {"passages": 2, "remarques_changeantes": []}
    assert rep["models"]["reranker"] == "faux-reranker" and rep["versions"]["app"]
    assert rep["mode"] == "rag"
    assert "bon tome @1 : 1/2" in out and "Référence : source essai" in out
    assert "Plancher de bruit : 0 remarque" in out


def test_banc_reponses_et_bruit(bench):
    llm = ScriptedLLM(GOOD, EVASIVE)  # 1er passage : bonne réponse ; 2e : deux esquives
    out, rep = bench("--answer", "--runs", "2", "--only", "Dubois-R1", llm=llm)
    first, second = (r["summary"] for r in rep["runs"])
    assert first["remarques"] == 1 and first["esquives"] == 0 and first["tome_cite"] == 1
    assert second["esquives"] == 1 and second["verdicts"]["INSUFFISANT"] == 1
    assert rep["noise"]["remarques_changeantes"] == ["Dubois-R1"]
    assert "esquives : 0/1" in out and "Plancher de bruit : 1 remarque(s)" in out
    assert rep["models"]["llm"] == "llm-scripte"


def test_banc_par_la_boucle(bench, tmp_path):
    """--agent : la boucle MRAe répond ; une esquive est rattrapée par sa passe
    élargie, comme dans le RAG seul."""
    loop = ROOT / "configs" / "profiles" / "mrae" / "loop.yaml"
    shutil.copy(loop, tmp_path / "profiles" / "ventes" / "loop.yaml")
    llm = ScriptedLLM(EVASIVE, GOOD)
    out, rep = bench("--agent", "--runs", "1", "--only", "Dubois-R1", llm=llm)
    summary = rep["runs"][0]["summary"]
    assert rep["mode"] == "boucle" and "(boucle)" in out
    assert summary["esquives"] == 0 and summary["verdicts"]["SUFFISANT"] == 1
    assert summary["tome_cite"] == 1 and len(llm.calls) == 2
    assert rep["models"]["llm"] == "llm-scripte"


def test_dossier_inconnu(bench, tmp_path):
    ref = tmp_path / "ref.jsonl"
    ref.write_text(
        json.dumps({"id": "X-R1", "dossier": "inconnu", "question": "q"}) + "\n", "utf-8"
    )
    with pytest.raises(ValueError, match="aucun document"):
        bench()


def test_resume_vide():
    assert summarize([], MRAE.metrics)["mrr_tome"] == 0
    assert noise([[]], MRAE.metrics) == {"passages": 1, "remarques_changeantes": []}
