"""DOCX, PPTX, XLSX, HTML, EML : fixtures générées, aucune donnée réelle."""

from pathlib import Path

import pytest

from luciole.ingestion.office import DocxParser, PptxParser, XlsxParser, _md_table
from luciole.ingestion.registry import default_registry
from luciole.ingestion.web_mail import EmlParser, HtmlParser


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def docx_file(tmp_path_factory) -> Path:
    import docx

    d = docx.Document()
    d.core_properties.title = "Compte rendu"
    d.core_properties.author = "Camille Martin"
    d.add_heading("Compte rendu de réunion", level=1)
    d.add_paragraph("Présents : Paul Dubois, Inès Leroy, Camille Martin.")
    d.add_heading("1. Besoins", level=2)
    d.add_paragraph("Retrouver une information en quelques minutes.")
    d.add_paragraph("Confidentialité des dossiers", style="List Bullet")
    d.add_paragraph("Budget contraint", style="List Bullet")
    d.add_heading("1.1 Détail", level=3)
    d.add_paragraph("Une journée par semaine perdue en recherche.")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Poste", "Coût"), ("Licence", "8 000 €"), ("Serveur", "4 000 €")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_heading("2. Suite", level=2)
    d.add_paragraph("Démonstration à la prochaine réunion.")
    p = tmp_path_factory.mktemp("docx") / "cr.docx"
    d.save(p)
    return p


@pytest.fixture(scope="module")
def pptx_file(tmp_path_factory) -> Path:
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # titre + contenu
    s1.shapes.title.text = "Luciole Pro"
    tf = s1.placeholders[1].text_frame
    tf.text = "Recherche locale dans vos dossiers"
    para = tf.add_paragraph()
    para.text = "Aucune donnée ne sort"
    para.level = 1
    s1.notes_slide.notes_text_frame.text = "Insister sur la confidentialité."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # titre seul
    s2.shapes.title.text = "Tarifs"
    shape = s2.shapes.add_table(3, 2, 0, 0, 100, 100)
    for r, (a, b) in enumerate([("Poste", "Prix"), ("Licence", "8 000 €"), ("Serveur", "4 000 €")]):
        shape.table.cell(r, 0).text, shape.table.cell(r, 1).text = a, b
    p = tmp_path_factory.mktemp("pptx") / "offre.pptx"
    prs.save(p)
    return p


@pytest.fixture(scope="module")
def xlsx_file(tmp_path_factory) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dossiers"
    ws.append(["Client", "Produit", "Commercial", "Issue"])
    for i in range(45):
        ws.append([f"Client {i}", "Luciole Pro", "Camille Martin", "gagne" if i % 3 else "perdu"])
    wb.create_sheet("Vide")
    ws3 = wb.create_sheet("Entête seule")
    ws3.append(["A", "B"])
    p = tmp_path_factory.mktemp("xlsx") / "suivi.xlsx"
    wb.save(p)
    return p


HTML = """<html><head><title>Fiche produit</title><style>p{}</style></head><body>
<nav><a href="/">Accueil</a></nav>
<h1>Luciole Pro</h1><p>Un assistant <b>local</b>.</p>
<h2>Atouts</h2><ul><li>Confidentialité</li><li>Citations exactes</li></ul>
<table><tr><th>Poste</th><th>Prix</th></tr><tr><td>Licence</td><td>8 000 €</td></tr></table>
<script>alert(1)</script><footer>© 2026</footer></body></html>"""

EML = """From: Camille Martin <camille@exemple.test>
To: Paul Dubois <paul@exemple.test>
Cc: ines@exemple.test
Subject: Proposition Luciole Pro
Date: Thu, 05 Mar 2026 09:12:00 +0100
Content-Type: text/plain; charset="utf-8"

Bonjour Paul,

Suite à la démonstration, voici la proposition chiffrée.

Bien à vous,
Camille
"""


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------
def test_md_table():
    assert _md_table([["a", "b|c"], ["1"]]) == "| a | b\\|c |\n|---|---|\n| 1 |  |"
    assert _md_table([]) == ""


def test_docx_structure(docx_file):
    doc = DocxParser().parse(docx_file)
    assert doc.metadata["title"] == "Compte rendu" and doc.metadata["author"] == "Camille Martin"
    types = [n.node_type for n in doc.nodes]
    assert types == [
        "title", "paragraph", "heading", "paragraph", "list_item", "list_item",
        "heading", "paragraph", "table", "heading", "paragraph",
    ]  # fmt: skip
    detail = doc.nodes[7]
    assert detail.section_path == ("Compte rendu de réunion", "1. Besoins", "1.1 Détail")
    table = doc.nodes[8]
    assert table.attributes["rows"] == 3 and "| Licence | 8 000 € |" in table.text
    assert table.section_path[-1] == "1.1 Détail", "le tableau garde sa place dans l'ordre"
    assert doc.nodes[10].section_path == ("Compte rendu de réunion", "2. Suite")
    assert doc.pages[0].quality == 1.0


def test_pptx_une_section_par_diapositive(pptx_file):
    doc = PptxParser().parse(pptx_file)
    assert doc.metadata["slide_count"] == 2 and len(doc.pages) == 2
    s1 = doc.nodes_on_page(1)
    assert s1[0].node_type == "heading" and s1[0].text == "Luciole Pro"
    assert [n.node_type for n in s1[1:]] == ["paragraph", "list_item", "paragraph"]
    assert s1[3].attributes == {"notes": True}
    s2 = doc.nodes_on_page(2)
    assert s2[0].text == "Tarifs" and s2[1].node_type == "table" and "Serveur" in s2[1].text
    assert all(n.section_path == ("Tarifs",) for n in s2[1:])


def test_xlsx_blocs_avec_entete_repete(xlsx_file):
    doc = XlsxParser(rows_per_block=20).parse(xlsx_file)
    assert doc.metadata["sheet_count"] == 3
    tables = [n for n in doc.nodes_on_page(1) if n.node_type == "table"]
    assert [t.attributes["rows"] for t in tables] == [21, 21, 6]
    assert all(t.text.startswith("| Client | Produit | Commercial | Issue |") for t in tables)
    assert tables[1].attributes["first_row"] == 22 and tables[2].attributes["last_row"] == 46
    assert "perdu" in tables[0].text
    assert doc.pages[1].method is None and doc.pages[1].quality == 0, "feuille vide"
    assert [n.attributes["rows"] for n in doc.nodes_on_page(3) if n.node_type == "table"] == [1]


def test_html(tmp_path):
    f = tmp_path / "fiche.html"
    f.write_text(HTML, encoding="utf-8")
    doc = HtmlParser().parse(f)
    assert doc.metadata["title"] == "Fiche produit"
    types = [n.node_type for n in doc.nodes]
    assert types == ["title", "paragraph", "heading", "list_item", "list_item", "table"]
    assert doc.nodes[1].text == "Un assistant local ."
    text = doc.to_text()
    assert "Accueil" not in text and "alert" not in text and "©" not in text
    assert "| Licence | 8 000 € |" in doc.nodes[5].text


def test_eml(tmp_path):
    f = tmp_path / "prop.eml"
    f.write_text(EML, encoding="utf-8")
    doc = EmlParser().parse(f)
    assert doc.metadata["subject"] == "Proposition Luciole Pro"
    assert doc.metadata["date"] == "2026-03-05"
    assert doc.metadata["from"].startswith("Camille Martin")
    assert doc.nodes[0].node_type == "title" and doc.nodes[0].text == "Proposition Luciole Pro"
    assert (
        doc.nodes[1].attributes == {"email_header": True}
        and "Cc : ines@exemple.test" in doc.nodes[1].text
    )
    assert [n.text for n in doc.nodes[2:]] == [
        "Bonjour Paul,",
        "Suite à la démonstration, voici la proposition chiffrée.",
        "Bien à vous, Camille",
    ]


def test_fichier_corrompu(tmp_path):
    f = tmp_path / "x.docx"
    f.write_bytes(b"pas un docx")
    with pytest.raises(ValueError, match="illisible"):
        DocxParser().parse(f)


def test_registre_complet(docx_file):
    reg = default_registry()
    for ext in (".docx", ".pptx", ".xlsx", ".html", ".eml", ".pdf", ".md", ".txt"):
        assert reg.supports(f"a{ext}"), ext
    assert reg.parse(docx_file).doc_type == "docx"
