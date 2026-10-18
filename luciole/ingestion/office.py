"""Formats bureautiques : DOCX, PPTX, XLSX, sur le modèle documentaire.

- DOCX : les styles de titre deviennent des ``heading`` (hiérarchie gardée),
  les listes des ``list_item``, les tableaux des ``table`` en Markdown, dans
  l'ordre réel du document (paragraphes et tableaux entrelacés).
- PPTX : une diapositive = une section (``heading`` = titre de la diapo),
  son corps, ses tableaux et ses notes ; ``page`` = numéro de diapositive.
- XLSX : une feuille = une section ; les lignes sont rendues en tableaux
  Markdown par lots avec l'en-tête répété, pour que chaque bloc reste
  lisible seul. Le magasin structuré (étape 2) relira les cellules.
"""

from __future__ import annotations

from pathlib import Path

from luciole.core.document import NodeBuilder, PageInfo, ParsedDocument
from luciole.ingestion.base import Parser


def _md_table(rows: list[list[str]]) -> str:
    """Tableau Markdown à partir de lignes de cellules (la première est l'en-tête)."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [
        [c.replace("|", "\\|").replace("\n", " ").strip() for c in r] + [""] * (width - len(r))
        for r in rows
    ]
    out = ["| " + " | ".join(norm[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(out)


def _single_page(chars: int, method: str) -> list[PageInfo]:
    return [
        PageInfo(
            number=1,
            method=method if chars else None,
            quality=1.0 if chars else 0.0,
            char_count=chars,
        )
    ]


class DocxParser(Parser):
    extensions = (".docx",)

    def parse(self, path: str | Path) -> ParsedDocument:
        import docx
        from docx.document import Document as _Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        try:
            d: _Document = docx.Document(str(path))
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"{path.name} : DOCX illisible ({e})") from e

        b = NodeBuilder()
        chars = 0
        title_seen = False
        body = d.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                p = Paragraph(child, d)
                text = " ".join(p.text.split())
                if not text:
                    continue
                chars += len(text)
                style = (p.style.name if p.style is not None else "") or ""
                lower = style.lower()
                if lower == "title" or (
                    lower.startswith("heading") and not title_seen and lower.endswith("1")
                ):
                    b.add("title", text, page=1, level=1, method="docx")
                    title_seen = True
                elif lower.startswith("heading") or lower.startswith("titre"):
                    digits = "".join(ch for ch in style if ch.isdigit())
                    level = int(digits) if digits else 2
                    b.add("heading", text, page=1, level=max(2, level + 1), method="docx")
                elif "list" in lower or p._p.pPr is not None and p._p.pPr.numPr is not None:
                    b.add("list_item", text, page=1, method="docx")
                elif lower.startswith("caption") or lower.startswith("légende"):
                    b.add("caption", text, page=1, method="docx")
                else:
                    b.add("paragraph", text, page=1, method="docx")
            elif child.tag == qn("w:tbl"):
                t = Table(child, d)
                rows = [[c.text for c in r.cells] for r in t.rows]
                md = _md_table(rows)
                if md:
                    chars += len(md)
                    b.add("table", md, page=1, method="docx", attributes={"rows": len(rows)})

        props = d.core_properties
        return ParsedDocument(
            file_name=path.name,
            source_path=str(path),
            doc_type="docx",
            metadata={
                "title": props.title or None,
                "author": props.author or None,
                "creation_date": props.created.isoformat() if props.created else None,
            },
            pages=_single_page(chars, "docx"),
            nodes=b.nodes(),
        )


class PptxParser(Parser):
    extensions = (".pptx",)

    def parse(self, path: str | Path) -> ParsedDocument:
        from pptx import Presentation
        from pptx.util import Emu  # noqa: F401 - import de contrôle de la dépendance

        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        try:
            prs = Presentation(str(path))
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"{path.name} : PPTX illisible ({e})") from e

        b = NodeBuilder()
        pages: list[PageInfo] = []
        for number, slide in enumerate(prs.slides, start=1):
            chars = 0
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            b.add("heading", title or f"Diapositive {number}", page=number, level=2, method="docx")
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[c.text for c in r.cells] for r in shape.table.rows]
                    md = _md_table(rows)
                    if md:
                        chars += len(md)
                        b.add(
                            "table", md, page=number, method="docx", attributes={"rows": len(rows)}
                        )
                    continue
                if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    text = " ".join(para.text.split())
                    if not text:
                        continue
                    chars += len(text)
                    kind = (
                        "list_item" if para.level > 0 or _looks_like_bullet(para) else "paragraph"
                    )
                    b.add(kind, text, page=number, method="docx")
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
                if notes:
                    chars += len(notes)
                    b.add(
                        "paragraph", notes, page=number, method="docx", attributes={"notes": True}
                    )
            pages.append(
                PageInfo(
                    number=number,
                    method="docx" if chars else None,
                    quality=1.0 if chars else 0.0,
                    char_count=chars,
                )
            )

        return ParsedDocument(
            file_name=path.name,
            source_path=str(path),
            doc_type="pptx",
            metadata={"title": prs.core_properties.title or None, "slide_count": len(pages)},
            pages=pages,
            nodes=b.nodes(),
        )


def _looks_like_bullet(para) -> bool:
    pPr = para._p.pPr
    if pPr is None:
        return False
    return any(child.tag.endswith(("buChar", "buAutoNum")) for child in pPr.iterchildren())


class XlsxParser(Parser):
    def __init__(self, rows_per_block: int = 20, max_rows_per_sheet: int = 20000) -> None:
        self.rows_per_block = rows_per_block
        self.max_rows_per_sheet = max_rows_per_sheet

    extensions = (".xlsx", ".xlsm")

    def parse(self, path: str | Path) -> ParsedDocument:
        import openpyxl

        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        try:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"{path.name} : XLSX illisible ({e})") from e

        b = NodeBuilder()
        pages: list[PageInfo] = []
        sheets_meta: list[str] = []
        try:
            for number, ws in enumerate(wb.worksheets, start=1):
                rows: list[list[str]] = []
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        rows.append(cells)
                    if len(rows) >= self.max_rows_per_sheet:
                        break
                b.add("heading", f"Feuille {ws.title}", page=number, level=2, method="docx")
                chars = 0
                if rows:
                    header, data = rows[0], rows[1:]
                    if not data:
                        md = _md_table([header])
                        chars += len(md)
                        b.add(
                            "table",
                            md,
                            page=number,
                            method="docx",
                            attributes={"rows": 1, "sheet": ws.title},
                        )
                    for start in range(0, len(data), self.rows_per_block):
                        block = data[start : start + self.rows_per_block]
                        md = _md_table([header, *block])
                        chars += len(md)
                        b.add(
                            "table",
                            md,
                            page=number,
                            method="docx",
                            attributes={
                                "rows": len(block) + 1,
                                "sheet": ws.title,
                                "first_row": start + 2,
                                "last_row": start + len(block) + 1,
                            },
                        )
                sheets_meta.append(f"{ws.title}:{len(rows)}")
                pages.append(
                    PageInfo(
                        number=number,
                        method="docx" if chars else None,
                        quality=1.0 if chars else 0.0,
                        char_count=chars,
                    )
                )
        finally:
            wb.close()

        return ParsedDocument(
            file_name=path.name,
            source_path=str(path),
            doc_type="xlsx",
            metadata={"sheet_count": len(pages), "sheets": ", ".join(sheets_meta)},
            pages=pages,
            nodes=b.nodes(),
        )
